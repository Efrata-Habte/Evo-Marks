"""Step 8: build docs/EvoMarks_slides.pptx (11 slides) from results/figures and the numbers in the report."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from config import ROOT

FIG = ROOT / "results" / "figures"
INK, MUTED, ACCENT = RGBColor(0x1F, 0x29, 0x37), RGBColor(0x6B, 0x72, 0x80), RGBColor(0x1D, 0x4E, 0xD8)
W, H = 13.333, 7.5

SLIDES = [
    ("Title", "EvoMarks: checking handwritten mark sheets with verified self-improvement and distillation",
     ["Read the marks and the total, check that they add up, flag the sheet. Offline, phone-sized.",
      "Both tracks: A (recursive learning with an independent verifier) and B (knowledge distillation)."],
     ["template/evomarks_template.png"]),
    ("Problem and design", "Arithmetic is the independent judge",
     ["5 marks + 2-digit total per sheet = 7 digits; ~5% of totals are genuinely wrong.",
      "Verifier: marks within maxima and sum equals total. It never sees confidence or labels.",
      "Splits: seed 150 (labeled), pool 3000 (unlabeled), dev 300, meta_dev 300, test 750 (unseen writers, read once).",
      "Fixed evaluator + sandbox + rollback; every version archived, accepted or rejected."], []),
    ("Method", "Three levels of improvement, then distillation",
     ["Level 0: self-training on sheets that pass the sum check; dev gate, rollback.",
      "Level 1: improver edits pipeline.py with AST operators; fixed evaluator accepts or rejects.",
      "Level 2: improver edits its own source; kept only if its utility on meta_dev is higher.",
      "Distillation: L = alpha CE + (1-alpha) T^2 KL; 25k-parameter student, T/alpha chosen on dev, 3 seeds."],
     []),
    ("Results: Level 0", "The sum check beats confidence (confirmation bias)",
     ["Seed-only teacher 0.887 -> 0.944 test digit accuracy.",
      "Pseudo-label precision 0.994 vs 0.968 for a per-digit confidence threshold; final dev 0.948 vs 0.928."],
     ["l0_verifier_vs_confidence.png", "l0_accuracy_by_round.png"]),
    ("Results: Level 1", "Mostly rejection; a small gain that does not survive full budget",
     ["20 candidates, 2 accepted, 18 rejected (all no_improvement).",
      "Noise floor (8 seeds): v000 dev 0.934 +- 0.005, 2.4x the gate margin 0.002.",
      "Paired by seed v016 - v000: +0.0105 +- 0.0039 (6/8 seeds) at evolution budget; at full budget test 0.9455 vs 0.9436 = within noise."],
     ["l1_evolution_timeline.png", "l1_edit_effectiveness.png"]),
    ("Results: Level 2", "A null result, and why",
     ["3 self-edits proposed, 0 accepted; Level 1 + 2 and Level 1 only are identical (0.9538 both).",
      "Utility is measured from a fixed parent, so PARENT_SELECTION edits cannot change it: a limit of the gate.",
      "The ablation therefore says nothing about whether recursion could work."],
     ["l2_improver_utility.png", "l2_ablation.png"]),
    ("Results: distillation", "3.2x smaller, 2.5-4x faster, 3.6 points less accurate",
     ["Hard-label 0.906 +- 0.002; distilled from v000 0.911 +- 0.006; from best 0.909 +- 0.001 (3 seeds).",
      "Distilled vs hard-label +0.3 points: no clear effect.",
      "Digit 9: teacher 21 points ahead; the average hides the tail."],
     ["kd_accuracy_vs_latency.png", "kd_per_digit_gap.png"]),
    ("Demo", "Reading a sheet",
     ["streamlit run app/streamlit_app.py  ->  Try it (upload a photo or a random synthetic sheet).",
      "Version vs previous version: colored code diff, metric deltas, per-digit bars.",
      "Backup: python demo.py photo.jpg"], []),
    ("Failure analysis", "What did not work, and why",
     ["Edits rarely helped: every operator type lowered dev accuracy on average; evolution is mostly rejection.",
      "The small-budget gain vanished at full budget: the loop optimised the cheap evaluator.",
      "Compensating errors pass the sum check (39 accepted sheets in the best run).",
      "Flag precision 0.14 at recall 1.0: most flags are misreads, not miscounts."],
     ["system_sheet_metrics.png"]),
    ("Interfaces between tracks", "Verified pseudo-labels feed the student",
     ["Teacher improved by verified feedback -> pseudo-labels (+ seed digits) -> student transfer set; teacher logits = soft targets.",
      "Noise in pseudo-labels passes into the student; the evolved teacher gave no better student than v000.",
      "Possible reverse link (not built): student as a cheap scorer inside the loop."], []),
    ("Limits", "What this does not show",
     ["Real handwriting: the real-photo test was not run; MNIST sheets are cleaner than phone photos.",
      "Small edit space, few seeds, single-training teachers, dev optimism after 20 candidates.",
      "Latency from one machine, one thread; the sandbox is not OS-level.",
      "Bounded, sandboxed recursion only: not open-ended self-improvement."], []),
]


def add_text(slide, left, top, width, height, text, size, bold=False, color=INK):
    """One text box with a single paragraph."""
    tf = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
    return tf


def add_picture(slide, path: Path, left, top, max_w, max_h) -> None:
    """Picture scaled to fit the box, aspect ratio kept."""
    w, h = Image.open(path).size
    scale = min(max_w / w, max_h / h)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(w * scale), Inches(h * scale))


def main() -> None:
    """Write the deck."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    for kicker, title, bullets, figs in SLIDES:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(s, 0.6, 0.35, 12, 0.4, kicker.upper(), 13, True, ACCENT)
        add_text(s, 0.6, 0.75, 12.1, 1.0, title, 28, True)
        paths = [FIG / f if not f.startswith("template") else ROOT / f for f in figs]
        text_w = 12.1 if not paths else 5.6
        tf = add_text(s, 0.6, 1.9, text_w, 4.8, bullets[0], 18)
        for b in bullets[1:]:
            tf.add_paragraph().text = b
        for p in tf.paragraphs:
            p.font.size, p.font.color.rgb, p.space_after = Pt(18), INK, Pt(10)
        if len(paths) == 1:
            add_picture(s, paths[0], 6.6, 1.9, 6.2, 4.9)
        elif len(paths) == 2:
            add_picture(s, paths[0], 6.5, 1.8, 6.4, 2.45)
            add_picture(s, paths[1], 6.5, 4.4, 6.4, 2.45)
    out = ROOT / "docs" / "EvoMarks_slides.pptx"
    prs.save(out)
    print(f"wrote {out} ({len(SLIDES)} slides)")


if __name__ == "__main__":
    main()
